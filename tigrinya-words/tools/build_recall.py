"""Free Recall answer content.

  python tools/build_recall.py draft                    # create or update tools/recall_review.xlsx
  python tools/build_recall.py build                    # write src/data/recall.json from reviewed rows
  python tools/build_recall.py build --include-drafts   # development only: draft rows too

Run build_data.py first. Reviewers edit recall_review.xlsx and set Status to "reviewed".
Only reviewed rows reach the app.
"""
import json
import re
import sys
import unicodedata
from collections import defaultdict
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / 'src' / 'data'
SHEET = ROOT / 'tools' / 'recall_review.xlsx'
RULE_VERSION = 1
CONTENT_TYPES = {'Noun', 'Adjective', 'Adverb', 'Number', 'Question word', 'Place name', 'Answer word'}
COLUMNS = ['Rank', 'Tigrinya', 'Pronunciation', 'Level', 'Word type', 'English prompt',
           'English to Tigrinya', 'Romanization alternatives', 'English answers', 'Status', 'Notes']
EDITABLE = ['English prompt', 'English to Tigrinya', 'Romanization alternatives',
            'English answers', 'Status', 'Notes']
APOSTROPHES = "'’‘ʼ`´"
DASHES = '-‐‑‒–—'


def first_meaning(meaning):
    return meaning.split(';')[0].strip()


def english_answers(meaning):
    """Draft answers: every comma- or semicolon-separated gloss, without notes in parentheses."""
    text = re.sub(r'\([^)]*\)', '', meaning)
    out = []
    for part in re.split(r'[;,]', text):
        p = ' '.join(part.split()).strip().lower()
        if p and p not in out:
            out.append(p)
    return out


def single_letter_spelling(pron):
    """Draft alternative: doubled consonants written once (al-lo -> alo). Reviewers approve or delete."""
    plain = pron.replace('-', '')
    single = re.sub(r'([bcdfghjklmnpqrstvwxyz])\1', r'\1', plain)
    return [single] if single != plain else []


def loose_key(text):
    """Same rules as the app's grader: what counts as 'the same spelling'."""
    s = unicodedata.normalize('NFC', text).strip().lower()
    for a in APOSTROPHES:
        s = s.replace(a, '')
    for d in DASHES:
        s = s.replace(d, '')
    s = re.sub(r'\s+', '', s)
    return s.replace('kh', 'h').replace('q', 'k')


def split_list(cell):
    return [x.strip() for x in str(cell or '').split(';') if x.strip()]


def draft():
    words = json.loads((DATA / 'words.json').read_text(encoding='utf-8'))
    lesson = [w for w in words if w['active'] and w['type'] in CONTENT_TYPES]
    shared = defaultdict(int)
    for w in words:
        if w['active']:
            shared[first_meaning(w['meaning']).lower()] += 1

    existing = {}
    if SHEET.exists():
        ws = load_workbook(SHEET).active
        head = [c.value for c in ws[1]]
        for r in ws.iter_rows(min_row=2, values_only=True):
            if r[0] is not None:
                existing[int(r[0])] = dict(zip(head, r))

    rows = []
    for w in sorted(lesson, key=lambda w: (w['level'], w['id'])):
        if w['id'] in existing:
            rows.append(existing[w['id']])
            continue
        rows.append({
            'Rank': w['id'], 'Tigrinya': w['word'], 'Pronunciation': w['pron'], 'Level': w['level'],
            'Word type': w['type'], 'English prompt': first_meaning(w['meaning']),
            'English to Tigrinya': 'yes' if shared[first_meaning(w['meaning']).lower()] == 1 else 'no',
            'Romanization alternatives': '; '.join(single_letter_spelling(w['pron'])),
            'English answers': '; '.join(english_answers(w['meaning'])),
            'Status': 'draft', 'Notes': '',
        })

    wb = Workbook()
    ws = wb.active
    ws.title = 'Recall answers'
    ws.append(COLUMNS)
    for r in rows:
        ws.append([r.get(c, '') for c in COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='1F3864')
        cell.alignment = Alignment(wrap_text=True, vertical='center')
    widths = [7, 14, 18, 7, 14, 26, 12, 26, 40, 11, 30]
    for i, wdt in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = wdt
    ws.freeze_panes = 'C2'
    ws.auto_filter.ref = ws.dimensions
    wb.save(SHEET)
    new = sum(1 for r in rows if r['Rank'] not in existing)
    print(f'{SHEET.name}: {len(rows)} rows ({new} new, {len(rows) - new} kept with your edits)')
    print(f'  English to Tigrinya = yes: {sum(1 for r in rows if r["English to Tigrinya"] == "yes")}')
    print(f'  rows with a drafted single-letter spelling: {sum(1 for r in rows if r["Romanization alternatives"])}')


def build(include_drafts):
    words = {w['id']: w for w in json.loads((DATA / 'words.json').read_text(encoding='utf-8'))}
    ws = load_workbook(SHEET).active
    head = [c.value for c in ws[1]]
    errors, warnings, items = [], [], []
    prompts = defaultdict(list)
    for r in ws.iter_rows(min_row=2, values_only=True):
        row = dict(zip(head, r))
        if row.get('Rank') is None:
            continue
        status = str(row.get('Status') or '').strip().lower()
        if status != 'reviewed' and not (include_drafts and status == 'draft'):
            continue
        wid = int(row['Rank'])
        w = words.get(wid)
        if not w or not w['active']:
            errors.append(f'rank {wid}: not a lesson word')
            continue
        answers = split_list(row.get('English answers'))
        if not answers:
            errors.append(f'{w["word"]} (rank {wid}): needs at least one English answer')
        alts, seen = [], {loose_key(w['pron'])}
        for a in split_list(row.get('Romanization alternatives')):
            k = loose_key(a)
            if not k:
                errors.append(f'{w["word"]}: an alternative is empty after normalizing')
            elif k in seen:
                warnings.append(f'{w["word"]}: alternative "{a}" is already accepted by the rules; removed')
            else:
                seen.add(k)
                alts.append(a)
        en_to_ti = str(row.get('English to Tigrinya') or '').strip().lower() == 'yes'
        prompt = str(row.get('English prompt') or '').strip()
        if en_to_ti:
            if not prompt:
                errors.append(f'{w["word"]}: English to Tigrinya needs an English prompt')
            prompts[prompt.lower()].append(w['word'])
        items.append({'wordId': wid, 'prompt': prompt, 'enToTi': en_to_ti,
                      'romanAlternatives': alts, 'englishAnswers': answers,
                      'status': status, 'ruleVersion': RULE_VERSION})

    for p, ws_ in prompts.items():
        if len(ws_) > 1:
            errors.append(f'English prompt "{p}" is used by {", ".join(ws_)}; make each prompt unique')

    # Collision audit (PRD section 6): different words that the grader would treat as the same spelling.
    keys = defaultdict(set)
    for it in items:
        w = words[it['wordId']]
        for form in [w['pron']] + it['romanAlternatives']:
            keys[loose_key(form)].add(w['word'])
    collisions = {k: v for k, v in keys.items() if len(v) > 1}

    if errors:
        print('Fix these before building:')
        for e in errors:
            print('  -', e)
        sys.exit(1)
    (DATA / 'recall.json').write_text(json.dumps(items, ensure_ascii=False, indent=1), encoding='utf-8')
    print(f'recall.json: {len(items)} words ({sum(1 for i in items if i["status"] == "reviewed")} reviewed)')
    print(f'  usable English to Tigrinya: {sum(1 for i in items if i["enToTi"])}')
    print(f'  usable Tigrinya to English: {len(items)}')
    for wmsg in warnings:
        print('  note:', wmsg)
    print(f'  spelling collisions to look at: {len(collisions)}')
    for k, v in sorted(collisions.items()):
        print(f'    "{k}": {" / ".join(sorted(v))}')


if __name__ == '__main__':
    if len(sys.argv) < 2 or sys.argv[1] not in ('draft', 'build'):
        print(__doc__)
        sys.exit(1)
    if sys.argv[1] == 'draft':
        draft()
    else:
        build('--include-drafts' in sys.argv)
