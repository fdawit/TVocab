"""Convert Tigrinya_Top1000_Learning.xlsx into words.json and units.json for the app.

Run from the project root:  python tools/build_data.py
Re-run it any time the spreadsheet changes.
"""
import json
import re
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / 'tools' / 'Tigrinya_Top1000_Learning.xlsx'
OUT_DIR = ROOT / 'src' / 'data'

UNIT_SIZE = 8          # target words per unit
MIN_UNIT = 4           # smaller leftovers join the previous unit; smaller topics go to Mixed Words
# Unit 1 is hand-picked so the first lesson is a real greeting (ranks from the Word List).
HELLO_UNIT = [319, 173, 776, 517, 61, 62, 709, 987]   # ሰላም ከመይ ደሓን እወ ኣነ እየ ሰናይ ብሰላም
# Topic order inside each level (also breaks ties when units are spread out).
TOPIC_ORDER = ['Feelings & Values', 'Questions', 'People & Family', 'Grammar Essentials',
               'Numbers & Amounts', 'Time & Calendar', 'Home & Everyday Life', 'Food & Drink',
               'Body & Health', 'Describing Words', 'School & Language', 'Places & Geography',
               'Nature & Animals', 'Actions', 'Faith & Religion', 'Sports & Games',
               'Culture & Celebrations', 'Work & Money', 'News & Government', 'Law & Courts', 'Names']

VARIANT_RE = re.compile(r'(short form of|softened|another spelling of|variant spelling of|'
                        r'variant of|older spelling of)\s+([\u1200-\u137F]+)', re.I)
BLANK = '＿＿＿'
GEEZ_PUNCT = '።፡፣?!.,"“”'


def tokens(sentence):
    return [t for t in sentence.split(' ') if t]


def bare(token):
    return token.strip(GEEZ_PUNCT)


def make_cloze(word, sentence):
    """Blank the word out of its example. Whole-token match first, then inside a token."""
    toks = tokens(sentence)
    for i, t in enumerate(toks):
        if bare(t) == word:
            toks[i] = t.replace(word, BLANK, 1)
            return ' '.join(toks)
    if word in sentence:
        return sentence.replace(word, BLANK, 1)
    return None


def addressee(english):
    m = re.search(r'\(to a (boy|girl)|\(to a group\)|\(said by a (boy|girl)\)', english)
    if not m:
        return None
    text = m.group(0)
    for key in ('boy', 'girl', 'group'):
        if key in text:
            return key
    return None


def main():
    ws = load_workbook(XLSX, read_only=True)['Word List']
    header = [c for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    col = {name: i for i, name in enumerate(header)}
    words = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[col['Rank']] is None:
            continue
        ex_ti = r[col['Example (Tigrinya)']] or ''
        ex_en = r[col['Example (English)']] or ''
        has_example = ex_ti not in ('', '—')
        words.append({
            'id': int(r[col['Rank']]),
            'word': r[col['Tigrinya Word']],
            'freq': int(r[col['Frequency']]),
            'pron': r[col['Pronunciation']],
            'meaning': r[col['English Meaning']],
            'type': r[col['Word Type']],
            'topic': r[col['Topic']],
            'level': int(str(r[col['Level']])[0]),
            'exTi': ex_ti if has_example else None,
            'exRom': r[col['Example (Romanized)']] if has_example else None,
            'exEn': ex_en if has_example else None,
            'note': r[col['Learning Note']] or '',
            'addressee': addressee(ex_en) if has_example else None,
        })
    by_word = {w['word']: w for w in words}
    by_id = {w['id']: w for w in words}

    # 1) Fold spelling variants and split forms into their main word.
    for w in words:
        w['mainId'] = None
        w['alsoWritten'] = []
        m = VARIANT_RE.search(w['meaning'])
        if m and m.group(2) in by_word and m.group(2) != w['word']:
            w['mainId'] = by_word[m.group(2)]['id']
    for w in words:
        if w['mainId']:
            by_id[w['mainId']]['alsoWritten'].append(w['word'])

    # 2) Exercise helpers.
    for w in words:
        w['cloze'] = make_cloze(w['word'], w['exTi']) if w['exTi'] else None
        toks = tokens(w['exTi']) if w['exTi'] else []
        w['canBuild'] = 3 <= len(toks) <= 6 and '...' not in w['exTi']
        w['active'] = w['level'] <= 3 and w['mainId'] is None and w['exTi'] is not None

    # 3) Units.
    active = [w for w in words if w['active']]
    hello = [by_id[i] for i in HELLO_UNIT]
    assert all(w['active'] and w['level'] == 1 for w in hello), 'Hello unit must be active Level 1 words'
    used = {w['id'] for w in hello}
    units = [{'id': 1, 'level': 1, 'topic': 'Greetings', 'title': 'Hello!',
              'wordIds': [w['id'] for w in hello]}]
    for level in (1, 2, 3):
        chunks_by_topic, mixed = {}, []
        for topic in TOPIC_ORDER:
            ws_ = sorted([w for w in active if w['level'] == level and w['topic'] == topic
                          and w['id'] not in used], key=lambda w: w['id'])
            if len(ws_) < MIN_UNIT:          # too few for its own unit: goes to a Mixed unit
                mixed.extend(ws_)
                continue
            chunks = [ws_[i:i + UNIT_SIZE] for i in range(0, len(ws_), UNIT_SIZE)]
            if len(chunks) > 1 and len(chunks[-1]) < MIN_UNIT:
                chunks[-2].extend(chunks.pop())
            chunks_by_topic[topic] = chunks
        if mixed:
            mixed.sort(key=lambda w: w['id'])
            chunks = [mixed[i:i + UNIT_SIZE] for i in range(0, len(mixed), UNIT_SIZE)]
            if len(chunks) > 1 and len(chunks[-1]) < MIN_UNIT:
                chunks[-2].extend(chunks.pop())
            chunks_by_topic['Mixed Words'] = chunks
        # Spread each topic's units evenly across the level so big topics
        # (like Grammar Essentials) don't end up in one long run at the end.
        order = TOPIC_ORDER + ['Mixed Words']
        slots = [(k / len(c), order.index(topic), topic, k, chunk)
                 for topic, c in chunks_by_topic.items() for k, chunk in enumerate(c)]
        for _, _, topic, k, chunk in sorted(slots, key=lambda s: s[:2]):
            many = len(chunks_by_topic[topic]) > 1
            units.append({'id': len(units) + 1, 'level': level, 'topic': topic,
                          'title': f'{topic} {k + 1}' if many else topic,
                          'wordIds': [w['id'] for w in chunk]})

    # 4) One say-it-at-home sentence per unit: the first short, complete sentence,
    #    or failing that the first example phrase in the unit.
    for u in units:
        full = [i for i in u['wordIds']
                if by_id[i]['exTi'][-1] in '።!?' and len(tokens(by_id[i]['exTi'])) <= 6]
        u['homeWordId'] = full[0] if full else u['wordIds'][0]

    # 5) Every active word must sit in exactly one unit.
    placed = [i for u in units for i in u['wordIds']]
    assert len(placed) == len(set(placed)) == len(active), (len(placed), len(set(placed)), len(active))

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    (OUT_DIR / 'words.json').write_text(json.dumps(words, ensure_ascii=False, indent=1), encoding='utf-8')
    (OUT_DIR / 'units.json').write_text(json.dumps(units, ensure_ascii=False, indent=1), encoding='utf-8')

    print(f'words.json: {len(words)} entries')
    print(f'  active lesson words (Levels 1-3, main forms): {len(active)}')
    print(f'  variants folded into a main word: {sum(1 for w in words if w["mainId"])}')
    print(f'  Level 4 (Word Bank only): {sum(1 for w in words if w["level"] == 4)}')
    print(f'  active words with a cloze sentence: {sum(1 for w in active if w["cloze"])}')
    print(f'  active words with a sentence-building exercise: {sum(1 for w in active if w["canBuild"])}')
    print(f'units.json: {len(units)} units ' +
          ', '.join(f'L{lv}: {sum(1 for u in units if u["level"] == lv)}' for lv in (1, 2, 3)))
    sizes = [len(u['wordIds']) for u in units]
    print(f'  words per unit: {min(sizes)}-{max(sizes)}')


if __name__ == '__main__':
    main()
